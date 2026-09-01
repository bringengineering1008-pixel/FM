from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


SOURCE = Path("blog/manuals/브링이슈_유튜브_콘텐츠_운영_매뉴얼_v1.0.docx")
OUTPUT = Path("blog/manuals/브링이슈_콘텐츠_운영_매뉴얼_v1.1.docx")


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def add_check(document: Document, text: str) -> None:
    document.add_paragraph(f"☐  {text}")


def main() -> None:
    document = Document(SOURCE)

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Version 1.0"):
            paragraph.text = "Version 1.1  |  2026.08.29  |  내부 운영용"
            break

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if "v1.0" in paragraph.text:
                paragraph.text = paragraph.text.replace("v1.0", "v1.1").replace("2026-08-21", "2026-08-29")

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    document.add_heading("부록 D. 네이버 블로그팀 공개 글 100편 반영 규칙", level=1)
    document.add_paragraph(
        "네이버 블로그팀 공식 블로그의 2025년 3월 28일부터 2026년 8월 27일까지 "
        "공개 글 100편을 본문 단위로 분석했다. 공식 공지·이벤트·소개 글 비중이 높은 표본이므로 "
        "문구를 복제하지 않고 모바일 가독성·시리즈 운영·시각자료 배치 원리만 브링이슈에 적용한다."
    )

    document.add_heading("D.1 관찰값과 적용 원칙", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["관찰 항목", "100편 중앙값·비율", "브링이슈 적용"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    rows = [
        ("본문 분량", "1,321자·57문단", "최소 글자 수로 강제하지 않고 질문 해결에 필요한 분량 사용"),
        ("모바일 문단", "문단 19자·60자 이하 90%", "1~3문장, 한 의미 묶음으로 분리"),
        ("이미지", "100% 사용·중앙값 3장", "대표·근거·행동·비교·예외 역할 수만큼 사용"),
        ("링크", "73% 사용·중앙값 2개", "공식 행동 링크와 다음 질문 관련 글만 연결"),
        ("구분선", "98% 사용·중앙값 2개", "기본 2개 안팎, 짧은 글 1개·긴 글 최대 4개"),
        ("이모지", "55% 사용·중앙값 2개", "본문 전체 1~3개, 제목은 필요할 때 0~1개"),
        ("표·인용구·영상", "10%·7%·1% 사용", "비교·판단·근거에 실제 필요할 때만 사용"),
        ("제목", "라벨 36%·숫자 71%·질문형 8%", "시리즈·연도·금액 등 실제 정보가 있을 때만 사용"),
    ]
    for item, observed, applied in rows:
        cells = table.add_row().cells
        cells[0].text = item
        cells[1].text = observed
        cells[2].text = applied

    document.add_heading("D.2 제목과 도입", level=2)
    add_bullet(document, "반복 가능한 후속 글이 있을 때만 8~12자의 짧은 시리즈 라벨을 제목 앞에 붙인다.")
    add_bullet(document, "숫자는 연도·월·회차·금액·조건처럼 실제 판단에 필요할 때만 사용한다.")
    add_bullet(document, "도입은 독자 장면 또는 질문 → 바로 필요한 결론 → 확인할 순서의 세 박자로 끝낸다.")
    add_bullet(document, "공식 블로그의 공지 말투인 ‘소개합니다·공개합니다’를 검색정보 글에 습관적으로 쓰지 않는다.")

    document.add_heading("D.3 문단과 네이버 서식", level=2)
    add_bullet(document, "일반 본문은 나눔고딕 19·줄간격 180%·가운데 정렬을 기본으로 한다.")
    add_bullet(document, "주요 소제목은 나눔고딕 24·굵게, 캡션은 나눔고딕 15·회색·가운데 정렬을 사용한다.")
    add_bullet(document, "긴 표·체크리스트·조건 비교는 왼쪽 정렬하고 표 본문은 나눔고딕 16을 사용한다.")
    add_bullet(document, "한 문단은 한 판단 또는 한 행동을 담당하며 1~3문장을 기본으로 한다.")
    add_bullet(document, "의미 묶음과 소제목·구분선·사진·표·인용구·링크 전후에는 빈 문단 1개만 둔다.")
    add_bullet(document, "빈 문단을 2개 이상 연속으로 넣어 큰 공백을 만들지 않는다.")
    add_bullet(document, "대상·상황·판단·행동이 바뀌는 지점에 빈 문단 또는 네이버 자체 구분선을 둔다.")
    add_bullet(document, "구분선은 기본 2개 안팎으로 사용하되 글의 실제 정보 블록에 따라 1~4개로 조정한다.")
    add_bullet(document, "이모지는 본문 전체 1~3개를 기본으로 하고 주요 소제목 또는 전환에만 사용한다.")
    add_bullet(document, "표는 동일 기준으로 둘 이상의 선택지를 비교할 때만 네이버 자체 표로 만든다.")
    add_bullet(document, "인용구는 최종 판단 한 문장을 강조할 때만 0~1개 사용한다.")
    add_bullet(document, "링크가 필요한 이유를 먼저 설명하고 바로 다음에 네이버 링크 카드를 넣으며 URL 목록을 붙이지 않는다.")

    document.add_heading("D.3.1 실제 관찰 이모티콘", level=3)
    add_bullet(document, "축하·새 소식: ✨ 26회, 🎉 24회, 🍀 22회, 🎁 13회, 🚀 10회")
    add_bullet(document, "일정·안내: 🗓️ 24회, 📌 13회, 📍 13회, 📝 13회, 📢 5회, 🔔 5회")
    add_bullet(document, "궁금증·확인: 👀 10회, 🤔 7회, 💡 5회, 🔍 4회")
    add_bullet(document, "주의·완료: ❗ 8회, ✔️ 7회, ⚠️ 6회, ✅ 6회")
    add_bullet(document, "금액·혜택: 💰 7회, 🍯 5회")
    document.add_paragraph(
        "브링이슈 생활정보 글은 📌·🔍·💡·⚠️·✅·💰·🗓️를 우선 사용한다. "
        "한 소제목에 최대 1개, 제목 0~1개, 글 전체 1~3개를 기본으로 하며 "
        "표·링크 문장·사진 캡션에는 넣지 않는다."
    )

    document.add_heading("D.3.2 네이버 캐릭터 스티커", level=3)
    add_bullet(document, "문자형 이모지와 캐릭터 스티커를 구분한다. 캐릭터 스티커는 스마트에디터 ONE의 스티커 컴포넌트에서 삽입한다.")
    add_bullet(document, "기본 제공·보유 스티커는 바로 사용하고, 추가 상품은 스티커 패널의 구매하기를 통해 네이버 OGQ 마켓에서 유료 또는 무료로 추가한다.")
    add_bullet(document, "도입·직접 실행·촬영 안내·체크 포인트·마무리 반응처럼 문맥이 맞는 곳에 적극 활용한다.")
    add_bullet(document, "개수의 상한이나 권장 수량을 두지 않는다. 각 스티커가 도입·설명·행동·확인·전환·마무리 중 분명한 역할을 가질 때 필요한 만큼 사용한다.")
    add_bullet(document, "많이 사용하는 것 자체를 목표로 삼지 않으며 같은 감정·같은 기능을 반복하거나 정보보다 스티커가 먼저 보이면 줄인다.")
    add_bullet(document, "정책·세금·피해·질병의 핵심 근거 옆에는 장난스러운 스티커를 두지 않는다.")
    add_bullet(document, "다른 블로그의 스티커를 캡처·복사하지 않고, 편집기 제공 또는 해당 계정이 정식 보유한 스티커만 사용한다.")

    document.add_heading("D.3.3 기존 글 내부 연결", level=3)
    add_bullet(document, "원고 작성 전에 공개된 브링이슈 기존 글의 제목·주제·URL을 검색한다.")
    add_bullet(document, "현재 문제의 다음 판단 또는 총론의 세부 행동으로 직접 이어지는 글만 1~2편 고른다.")
    add_bullet(document, "예약·비공개 글은 연결하지 않고 실제 공개 URL을 다시 열어 확인한 뒤 사용한다.")
    add_bullet(document, "연결 이유를 한 문장으로 쓴 뒤 바로 아래에 네이버 링크 카드를 넣고 URL 목록을 붙이지 않는다.")

    document.add_heading("D.4 시각자료 역할표", level=2)
    image_table = document.add_table(rows=1, cols=2)
    image_table.style = "Table Grid"
    image_table.rows[0].cells[0].text = "역할"
    image_table.rows[0].cells[1].text = "사용 기준"
    image_rows = [
        ("대표", "글이 해결하는 문제를 한눈에 설명"),
        ("근거", "공식 사이트·앱에서 실제로 확인되는 위치 증명"),
        ("행동", "다음에 누를 메뉴나 입력 항목 안내"),
        ("비교", "선택지 차이를 같은 기준으로 정리"),
        ("예외", "놓치기 쉬운 조건과 위험을 표시"),
    ]
    for role, rule in image_rows:
        cells = image_table.add_row().cells
        cells[0].text = role
        cells[1].text = rule
    document.add_paragraph(
        "이미지 수는 고정하지 않는다. 공식 화면은 잘리지 않고 모바일에서 글자가 읽혀야 한다. "
        "AI·Google Flow 이미지는 대표·설명·비교 보조에는 사용할 수 있지만 공식 사이트·앱 화면을 대체하지 않는다."
    )

    document.add_heading("D.5 반복 시리즈", level=2)
    add_bullet(document, "한 주제는 총론 1편과 세부 질문 3~6편으로 확장한다.")
    add_bullet(document, "시리즈 라벨과 대표색은 유지하되 제목·도입·결론 문장을 복사하지 않는다.")
    add_bullet(document, "후속 글은 앞 글의 요약이 아니라 독자의 다음 질문 하나를 완결한다.")
    add_bullet(document, "관련 글은 현재 글을 읽은 뒤 자연스럽게 생기는 다음 질문 1~2편만 연결한다.")

    document.add_heading("D.6 적용 금지", level=2)
    add_bullet(document, "이벤트 글의 많은 댓글을 일반 정보글의 성과 공식으로 해석하지 않는다.")
    add_bullet(document, "모든 제목에 대괄호·숫자·이모지를 동시에 넣지 않는다.")
    add_bullet(document, "이미지 3장·구분선 2개·이모지 2개를 기계적인 정답으로 만들지 않는다.")
    add_bullet(document, "영상 캡처나 AI 이미지를 공식 근거 화면 대신 사용하지 않는다.")

    document.add_heading("D.7 발행 전 추가 체크", level=2)
    add_check(document, "제목 라벨·숫자·이모지는 실제 정보와 시리즈 목적이 있을 때만 썼다.")
    add_check(document, "캐릭터 스티커는 네이버 편집기 제공 또는 해당 계정이 정식 보유한 항목만 사용했고 도입·단계·체크·마무리의 역할과 문맥이 맞는다.")
    add_check(document, "공개된 기존 브링이슈 글을 검색하고, 직접 이어지는 글이 있으면 공개 URL을 확인해 링크 카드 1~2개로 연결했다.")
    add_check(document, "도입 세 문단 안에서 독자 장면·결론·확인 순서가 보인다.")
    add_check(document, "각 이미지에 대표·근거·행동·비교·예외 중 하나의 역할이 있다.")
    add_check(document, "구분선·이모지·표를 관찰 중앙값에 맞추려고 억지로 넣지 않았다.")
    add_check(document, "본문 19·소제목 24·캡션 15·표 16과 줄간격 180%를 실제 편집기에서 확인했다.")
    add_check(document, "의미 묶음과 각 편집 요소 전후의 빈 문단이 1개이며 연속 빈 문단이 없다.")
    add_check(document, "공식 링크와 관련 글을 설명 문장 뒤 네이버 링크 카드로 넣었다.")
    add_check(document, "실제 관찰 목록에서 주제 역할에 맞는 이모티콘을 골랐고 같은 문자를 반복하지 않았다.")
    add_check(document, "공식 화면이 잘리지 않고 모바일에서 글자를 읽을 수 있다.")
    add_check(document, "시리즈 후속 글은 이전 글을 복사하지 않고 다음 질문을 해결한다.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
