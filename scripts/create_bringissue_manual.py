from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"C:\Users\user\.codex\worktrees\7129\마케팅\blog\manuals\브링이슈_유튜브_콘텐츠_운영_매뉴얼_v1.0.docx")

BLUE = "2E74B5"
DARK_BLUE = "18324A"
MID_BLUE = "4F81BD"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E3F6D3"
PALE_YELLOW = "FFF4CC"
PALE_RED = "FCE8E6"
GRAY = "667085"
LIGHT_GRAY = "F4F6F8"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="CBD5E1", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:cantSplit")
    tr_pr.append(tag)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=False, color=BLACK, italic=False, underline=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.color.rgb = RGBColor.from_string(color)


def shade_run(run, fill):
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    r_pr.append(shd)


def add_custom_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def make_abstract(fmt, text, abs_id, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        first_num_index = next((i for i, child in enumerate(numbering) if child.tag == qn("w:num")), len(numbering))
        numbering.insert(first_num_index, abstract)

    def make_num(abs_id, num_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), str(abs_id))
        num.append(abstract_id)
        lvl_override = OxmlElement("w:lvlOverride")
        lvl_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        lvl_override.append(start_override)
        num.append(lvl_override)
        numbering.append(num)

    make_abstract("bullet", "•", next_abs)
    make_abstract("decimal", "%1.", next_abs + 1)
    make_num(next_abs, next_num)
    decimal_nums = []
    for offset in range(1, 9):
        make_num(next_abs + 1, next_num + offset)
        decimal_nums.append(next_num + offset)
    return next_num, decimal_nums


def set_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("BRINGISSUE  |  YouTube Content Operations Manual")
    set_run_font(hr, 8.5, True, GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("브링이슈 내부 운영문서  ·  v1.0  ·  2026-08-21")
    set_run_font(fr, 8, False, GRAY)


def add_title_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(86)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BRINGISSUE")
    set_run_font(r, 12, True, MID_BLUE)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run("유튜브 콘텐츠\n운영 매뉴얼")
    set_run_font(r2, 30, True, DARK_BLUE)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(28)
    r3 = p3.add_run("흥하는 영상을 발견하고, 새 해석을 만들고,\n검색·홈피드 유입과 검증된 수익화로 연결하는 실전 SOP")
    set_run_font(r3, 13, False, GRAY)
    add_callout(doc, "핵심 운영 공식", "영상 복제 X  ·  원문 파악 O  ·  에디터 해설 O  ·  장면 설계 O  ·  검증된 상품만 연결", PALE_GREEN)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(54)
    r4 = p4.add_run("Version 1.0  |  2026.08.21  |  내부 운영용")
    set_run_font(r4, 9.5, True, GRAY)
    doc.add_page_break()


def add_callout(doc, label, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_run_font(r, 10.5, True, DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, 10.5, False, BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc, text, bullet_num):
    p = doc.add_paragraph()
    set_num(p, bullet_num)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_number(doc, text, decimal_num):
    p = doc.add_paragraph()
    set_num(p, decimal_num)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_para(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 11, False, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, 11, False, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, text in enumerate(headers):
        set_cell_shading(hdr.cells[i], header_fill)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, font_size, True, DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            set_run_font(r, font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_checklist(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("☐  ")
        set_run_font(r, 11, True, MID_BLUE)
        r2 = p.add_run(item)
        set_run_font(r2, 10.7)


def add_form_line(doc, label, hint=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, 10.5, True, DARK_BLUE)
    r2 = p.add_run(hint if hint else "____________________________________________________________")
    set_run_font(r2, 10, False, GRAY)


def build():
    doc = Document()
    configure_doc(doc)
    bullet_num, decimal_nums = add_custom_numbering(doc)
    add_title_cover(doc)

    add_heading(doc, "문서 사용법", 1)
    add_callout(doc, "이 매뉴얼의 목적", "누가 작업하더라도 같은 판단 순서와 같은 품질 기준으로 ‘브링이슈’ 글을 기획·제작·발행·회고하도록 만드는 운영 기준서입니다.")
    add_para(doc, "이 문서는 읽고 끝내는 자료가 아닙니다. 영상 한 편을 선택할 때마다 2장부터 순서대로 실행하고, 부록의 빈 양식을 복사해 작업 원장으로 사용합니다.")
    add_table(doc, ["사용 시점", "열어볼 장"], [
        ("영상 후보를 고를 때", "2장 후보 선별 + 3장 원문 분석"),
        ("제목·각도를 정할 때", "4장 정보 기획 + 5장 제목 설계"),
        ("캡처할 때", "6장 장면·캡처 설계"),
        ("본문을 쓸 때", "7장 고정 글 구조 + 8장 네이버 서식"),
        ("업로드 직전", "9장 발행 절차 + 10장 권리·사실 검수"),
        ("발행 후", "12장 성과 측정과 개선"),
    ], [2700, 6660])
    add_para(doc, "문서 상태: v1.0. 첫 10개 글은 학습 구간으로 운영하며, 실제 유입과 전환 데이터를 근거로 분기별 개정합니다.", italic=True)

    add_heading(doc, "1. 브링이슈 운영 원칙", 1)
    add_heading(doc, "1.1 브랜드 정의", 2)
    add_para(doc, "브링이슈는 인기 유튜브 영상을 그대로 옮기는 요약 블로그가 아닙니다. 사람들이 반응하는 사건·선택·관계·돈의 포인트를 찾아, 원본을 본 에디터의 해설과 추가 맥락으로 다시 읽히게 만드는 콘텐츠 브랜드입니다.")
    add_callout(doc, "한 줄 정의", "흥하는 유튜브를 발견한다 → 전체 맥락을 파악한다 → 한 가지 질문으로 재구성한다 → 독자에게 새로운 해석을 준다 → 검증된 경우에만 구매·서비스로 연결한다.", PALE_GREEN)
    add_heading(doc, "1.2 콘텐츠와 수익의 비율", 2)
    add_table(doc, ["영역", "권장 비중", "역할"], [
        ("이슈·스토리 유입", "70%", "조회, 체류, 댓글, 재방문을 만드는 본편"),
        ("구매·제휴 확장", "30%", "공식적으로 확인된 제품·서비스만 별도 후속 글로 연결"),
    ], [2100, 1500, 5760])
    add_bullet(doc, "본편에 억지 상품을 끼워 넣지 않는다.", bullet_num)
    add_bullet(doc, "제품명이 공식 영상·브랜드·착장 정보로 확인되지 않으면 ‘비슷한 제품’이라도 단정하지 않는다.", bullet_num)
    add_bullet(doc, "제휴 링크가 있으면 링크 전에 광고·제휴 사실을 명확히 표시한다.", bullet_num)
    add_heading(doc, "1.3 반드시 지키는 편집 원칙", 2)
    add_checklist(doc, [
        "전체 대사와 영상 흐름을 먼저 파악한 뒤 기획한다.",
        "한 글에는 한 독자, 한 질문, 한 약속, 한 행동만 둔다.",
        "대사를 길게 복사하지 않고 필요한 짧은 표현만 맥락과 함께 쓴다.",
        "캡처는 본문을 대신하는 장식이 아니라 주장과 감정을 증명하는 장면으로 사용한다.",
        "확인되지 않은 사생활·재산·의료·관계 추측은 쓰지 않는다.",
        "홈피드와 검색 노출을 보장하지 않고 실제 유입 경로로 학습한다.",
    ])

    add_heading(doc, "2. 영상 후보 선별 시스템", 1)
    add_heading(doc, "2.1 영상 접수 카드", 2)
    for label, hint in [
        ("영상 URL", ""), ("채널 / 영상 제목", ""), ("업로드 시각 / 영상 길이", ""),
        ("확인 시점 조회수·좋아요·댓글", ""), ("등장인물", ""),
        ("영상 유형", "인터뷰 / 브이로그 / 사건 / 제품 / 돈 / 관계 / 기타"),
        ("예상 독자", "누가 왜 클릭할 것인가?"), ("한 줄 사건", "영상에서 실제로 벌어진 변화는 무엇인가?"),
    ]:
        add_form_line(doc, label, hint)
    add_heading(doc, "2.2 100점 평가표", 2)
    add_table(doc, ["평가 항목", "배점", "판단 질문"], [
        ("현재 관심", "20", "지금 사람들이 실제로 보고·검색·공유하고 있는가?"),
        ("클릭 가능성", "15", "제목 한 줄로 사건의 변화나 의외성을 설명할 수 있는가?"),
        ("검색 의도", "15", "인물·영상·사건을 궁금해하는 독자의 질문이 분명한가?"),
        ("독자 가치", "15", "원본 링크만 보는 것보다 우리 글을 읽을 이유가 있는가?"),
        ("근거·이미지", "10", "대사, 공식 정보, 필요한 장면을 확인할 수 있는가?"),
        ("차별성", "10", "단순 요약이 아닌 해석·구조·비교가 가능한가?"),
        ("확장성", "5", "후속 검색 글이나 검증된 상품 글로 이어질 수 있는가?"),
        ("포트폴리오 보정", "10", "최근 발행 비중에서 필요한 콘텐츠인가?"),
    ], [2100, 900, 6360], font_size=9.2)
    add_callout(doc, "작성 기준", "75점 이상만 원고 제작에 들어갑니다. 단, 사실성·권리·개인정보·독자 안전 중 하나라도 실패하면 총점과 관계없이 제외합니다.", PALE_YELLOW)
    add_heading(doc, "2.3 즉시 제외할 후보", 2)
    add_bullet(doc, "확인되지 않은 열애·불화·재산·건강·범죄 피해를 자극적으로 소비해야만 성립하는 주제", bullet_num)
    add_bullet(doc, "원본 영상을 거의 그대로 재배포해야만 설명할 수 있는 주제", bullet_num)
    add_bullet(doc, "상품 연결이 영상 내용과 무관하거나 제품 확인 근거가 없는 주제", bullet_num)
    add_bullet(doc, "제목의 결론을 본문에서 확인할 수 없는 주제", bullet_num)

    add_heading(doc, "3. 전체 대사·영상 분석 프로토콜", 1)
    add_callout(doc, "선행 원칙", "캡처보다 먼저 대사를 파악합니다. 무작위 캡처를 금지하고, ‘어떤 정보를 알려줄지’가 정해진 다음에만 장면표를 만듭니다.", PALE_RED)
    add_heading(doc, "3.1 1차 시청: 사실 지도 만들기", 2)
    decimal_num = decimal_nums[0]
    for text in [
        "영상 전체를 처음부터 끝까지 보고 등장인물·장소·시간·행동을 기록한다.",
        "조회수·좋아요·댓글 등 변동 수치는 확인 시각과 함께 별도 기록한다.",
        "확인된 사실, 영상 속 주장, 에디터 해석을 서로 다른 칸으로 분리한다.",
        "개인정보, 화면 속 전화번호·계정·차량번호 등 가림 대상을 표시한다.",
    ]:
        add_number(doc, text, decimal_num)
    doc.add_page_break()
    add_heading(doc, "3.2 2차 시청: 감정과 전환 지도", 2)
    add_table(doc, ["구간", "찾아야 할 것", "기록 예시"], [
        ("시작", "왜 이 영상을 봐야 하는가", "팬심·목표·갈등의 출발"),
        ("배경", "인물의 이유와 맥락", "과거 선택, 취향, 관계"),
        ("복선", "뒤의 결과를 기대하게 하는 정보", "장소 이동, 질문, 단서"),
        ("전환", "상황이 달라지는 순간", "예상 밖 인물·답변·발견"),
        ("클라이맥스", "가장 큰 보상", "만남, 제안, 계약, 결과"),
        ("여운", "독자가 댓글로 말하고 싶은 감정", "놀람, 공감, 부러움, 의문"),
    ], [1300, 3400, 4660], font_size=9.2)
    add_heading(doc, "3.3 독자 질문으로 변환", 2)
    add_para(doc, "타임라인을 그대로 요약하지 말고 아래 질문 3~6개로 바꿉니다. 이 질문이 그대로 본문 소제목과 정보 구조가 됩니다.")
    add_bullet(doc, "처음에는 왜 이 행동을 시작했나?", bullet_num)
    add_bullet(doc, "중간에 어떤 단서가 상황을 바꿨나?", bullet_num)
    add_bullet(doc, "결국 실제로 무슨 일이 벌어졌나?", bullet_num)
    add_bullet(doc, "사람들이 이 장면에 반응한 이유는 무엇인가?", bullet_num)
    add_bullet(doc, "원본 영상에서 놓치기 쉬운 선택이나 태도는 무엇인가?", bullet_num)
    add_bullet(doc, "이 사건에서 검증 가능한 후속 정보는 무엇인가?", bullet_num)

    add_heading(doc, "4. 정보 기획과 글의 각도", 1)
    add_heading(doc, "4.1 한 글의 기획 문장", 2)
    add_callout(doc, "기획 문장 공식", "[누가] [어떤 선택·행동]을 했고, [예상 밖 결과]로 이어진 과정을 보여주며, 독자에게 [새로운 해석]을 준다.", PALE_GREEN)
    add_heading(doc, "4.2 각도 선택 우선순위", 2)
    add_table(doc, ["각도", "언제 선택하나", "핵심 독자 가치"], [
        ("사건·결과형", "결말이 강하고 과정이 궁금할 때", "결과까지 가는 흐름"),
        ("인물·선택형", "성격과 결정이 반응을 만들었을 때", "사람을 이해하는 해설"),
        ("돈·계약형", "가격, 수익, 계약, 사업 판단이 있을 때", "의사결정 기준"),
        ("제품·정보형", "공식적으로 제품이나 방법을 확인할 수 있을 때", "구매·실행 정보"),
    ], [1800, 3600, 3960], font_size=9.3)
    add_heading(doc, "4.3 본편과 수익 글 분리 규칙", 2)
    add_para(doc, "본편은 사건과 감정의 완결성을 우선합니다. 제품·서비스는 공식 근거를 별도 확인한 뒤 후속 글로 분리하는 것을 기본값으로 둡니다.")
    add_table(doc, ["본편", "후속 수익 글"], [
        ("왜 이런 일이 벌어졌는지", "공식 제품명·가격·구매처 확인"),
        ("어떤 장면에서 감정이 바뀌었는지", "비교 기준·대체재·주의점"),
        ("브링이슈만의 인물·선택 해설", "광고 고지 후 단일 구매 행동"),
    ], [4680, 4680], header_fill=PALE_GREEN)

    add_heading(doc, "5. 자극적이되 낚시가 아닌 제목 설계", 1)
    add_heading(doc, "5.1 제목의 3요소", 2)
    add_bullet(doc, "주체: 누가 등장하는지 첫눈에 알 수 있어야 한다.", bullet_num)
    add_bullet(doc, "사건: 무엇을 했거나 어떤 변화가 있었는지 보여준다.", bullet_num)
    add_bullet(doc, "빈칸: 결론 일부는 남기되 본문에서 반드시 회수한다.", bullet_num)
    add_heading(doc, "5.2 반복 사용 가능한 공식", 2)
    add_table(doc, ["공식", "구조", "주의점"], [
        ("사건→결과", "[인물]이 [행동]했더니, 결국 [결과]", "결과를 과장하지 않는다"),
        ("과정→궁금증", "[목표] 때문에 [장소]까지 간 이유", "본문이 실제 이유를 설명해야 한다"),
        ("짧은 말→전환", "‘[짧은 표현]’ 그 한마디 뒤 벌어진 일", "짧은 인용만 사용한다"),
        ("통념→반전", "다들 [통념]이라 봤지만 영상에서 보인 건 달랐다", "우리 해석임을 분명히 한다"),
    ], [1700, 3900, 3760], font_size=9.2)
    add_heading(doc, "5.3 제목 품질 체크", 2)
    add_checklist(doc, [
        "제목의 핵심 명사와 사건이 도입 5문단 안에 등장한다.",
        "영상에 없는 감정·의도·관계를 단정하지 않는다.",
        "‘충격’, ‘경악’, ‘난리’ 같은 추상 과장어보다 실제 사건을 쓴다.",
        "모바일 화면에서 앞부분만 보여도 인물과 사건을 이해할 수 있다.",
        "추천 제목 1개 외에 대안 2개를 남겨 발행 전 비교한다.",
    ])

    add_heading(doc, "6. 장면·캡처·클립 설계", 1)
    add_heading(doc, "6.1 장면표는 글의 증거표", 2)
    add_para(doc, "캡처 수를 먼저 정하지 않습니다. 도입·복선·전환·클라이맥스·여운을 증명하는 장면만 선택하며, 보통 7~10장이면 충분합니다.")
    add_table(doc, ["번호", "타임코드", "본문 역할", "화면 내용", "캡션", "가림/주의"], [
        ("1", "", "도입", "", "", ""),
        ("2", "", "배경", "", "", ""),
        ("3", "", "복선", "", "", ""),
        ("4", "", "전환", "", "", ""),
        ("5", "", "대표", "", "", ""),
        ("6", "", "클라이맥스", "", "", ""),
        ("7", "", "여운", "", "", ""),
    ], [650, 1050, 1250, 2200, 2600, 1610], font_size=8.3)
    add_heading(doc, "6.2 깨끗한 캡처 기준", 2)
    add_checklist(doc, [
        "16:9 영상 프레임만 남기고 YouTube 제목·버튼·재생바·댓글·확장프로그램 UI를 제거한다.",
        "원본 영상 안에 포함된 채널 워터마크·제작자 표시는 지우지 않는다.",
        "상황 이해에 필요한 영상 자막은 남기되, 의미 없는 플레이어 자막은 최소화한다.",
        "전화번호·메신저 ID·차량번호·문서·위치 등 개인정보를 가린다.",
        "얼굴을 과도하게 확대하거나 맥락을 왜곡하는 크롭을 피한다.",
        "대표 이미지는 가장 큰 결과가 드러나는 한 장을 선택한다.",
    ])
    add_heading(doc, "6.3 제미나이 매거진 일러스트", 2)
    add_para(doc, "문장 전용 정보 카드는 사용하지 않습니다. 공개 영상에 없는 배경 상황·감정·맥락은 제미나이로 제작한 세련된 한국 매거진 에디토리얼 일러스트로 시각화합니다. 실제 캡처는 사실의 증거, AI 이미지는 이해를 돕는 해설이라는 역할을 분리합니다.")
    add_table(doc, ["구분", "실제 캡처", "제미나이 일러스트"], [
        ("역할", "인물·발언·사건·결과를 증명", "배경·감정·관계·해석을 설명"),
        ("배치", "대표·도입·전환·클라이맥스", "해설·맥락 문단 뒤"),
        ("기준", "출처와 타임코드 확인", "AI 고지 캡션과 비식별 표현"),
        ("금지", "YouTube UX/UI와 왜곡 크롭", "실제 보도사진처럼 보이는 합성과 얼굴 복제"),
    ], [1400, 3930, 4030], font_size=8.9)
    add_checklist(doc, [
        "16:9 가로형, 1200×675px 이상으로 제작한다.",
        "실존 인물의 얼굴을 그대로 복제하지 않고 인상·의상·분위기만 참고한 비식별 인물로 표현한다.",
        "이미지 안에 제목·긴 문장·말풍선·로고·워터마크·유튜브 UI를 넣지 않는다.",
        "AI 이미지를 실제 사건의 증거처럼 배치하지 않고 연속 두 장 이상 사용하지 않는다.",
        "모든 AI 이미지 아래에 ‘AI로 제작한 이해용 이미지입니다.’라고 표시한다.",
    ])
    add_callout(doc, "제미나이 공통 프롬프트", "한국 온라인 매거진에 들어갈 16:9 가로형 에디토리얼 일러스트.\n\n장면: [누가 / 어디서 / 무엇을 하는지]\n핵심 감정: [놀람 / 배려 / 긴장 / 반전 / 따뜻함]\n주요 소품: [문단을 설명하는 물건 1~3개]\n구도: [와이드숏 / 미디엄숏 / 손과 소품 중심 / 공간 중심]\n색감: 밝고 세련된 크림·옐로 / 블루 / 피치 계열.\n스타일: 고급 라이프스타일 매거진 삽화, 자연스러운 빛, 충분한 여백, 현실적인 공간감, 섬세한 질감.\n\n실제 보도사진처럼 만들지 말 것. 실존 인물의 얼굴을 그대로 복제하지 말 것. 글자, 자막, 로고, 워터마크, 말풍선, 유튜브 UI를 넣지 말 것. 해상도 1200×675 이상.", PALE_YELLOW)
    doc.add_page_break()
    add_heading(doc, "6.4 짧은 녹화·GIF 기준", 2)
    add_bullet(doc, "정지 화면으로 의미가 전달되지 않는 표정·움직임에만 사용한다.", bullet_num)
    add_bullet(doc, "한 클립은 대체로 3~7초, 연속 장면 재배포는 피한다.", bullet_num)
    add_bullet(doc, "원본을 대체하지 않도록 해설 문단과 가까이 배치한다.", bullet_num)
    add_bullet(doc, "원본 채널명·영상 제목·링크를 본문 끝 출처에 명시한다.", bullet_num)

    add_heading(doc, "7. 브링이슈 고정 글 구조", 1)
    add_callout(doc, "핵심 리듬", "결과를 살짝 보여준다 → 이유를 만든다 → 복선을 쌓는다 → 전환을 보여준다 → 브링이슈의 해석으로 마무리한다.", PALE_GREEN)
    sections = [
        ("1) 제목", "인물 + 실제 사건 + 궁금증. 본문에서 반드시 회수할 약속만 쓴다."),
        ("2) 도입 3~5문단", "가장 강한 결과를 먼저 암시하고, 독자가 궁금해할 질문을 한 번만 던진다."),
        ("3) 핵심 인용구", "이 글의 판단 또는 전환을 짧게 정리한다. 핵심 결론 한 문장만 밑줄 처리한다."),
        ("4) 대표 이미지", "글의 가장 큰 결과가 보이는 장면. 구체적인 캡션을 붙인다."),
        ("5) 배경·복선", "왜 시작했는지, 어떤 단서가 있었는지. 영상 순서를 그대로 늘어놓지 않는다."),
        ("6) 구분선", "배경에서 전환으로 넘어가는 한 번의 호흡. 문자선이 아니라 네이버 구분선 컴포넌트를 쓴다."),
        ("7) 반전·전환", "예상이 바뀐 순간을 캡처와 함께 설명한다."),
        ("8) 클라이맥스", "제안·만남·결과 등 제목의 약속을 회수한다."),
        ("9) 제미나이 이해 이미지", "공개 장면에 없는 배경·감정·맥락만 시각화하고 AI 고지 캡션을 붙인다."),
        ("10) 브링이슈 해설", "왜 사람들이 반응했는지, 어떤 선택이 결과를 만들었는지 2~4문단으로 해석한다."),
        ("11) 댓글 질문", "정답형 질문이 아니라 독자의 경험·취향을 묻는 질문 하나만 둔다."),
        ("12) 출처", "원본 채널명, 영상 제목, 링크, 확인일을 간결하게 적는다."),
    ]
    for label, body in sections:
        add_para(doc, f"{label}  {body}", bold_prefix=label)
    add_heading(doc, "7.1 복사해서 쓰는 본문 골격", 2)
    add_callout(doc, "[최종 제목]", "[결과를 암시하는 짧은 도입]\n\n[왜 이런 일이 벌어졌을까?]\n\n> [핵심 판단 — 밑줄 1회]\n\n[대표 이미지 + 캡션]\n\n[이모지 소제목] 시작과 배경\n[사실 + 실제 장면 + 해설]\n\n[제미나이 이해 이미지]\nAI로 제작한 이해용 이미지입니다.\n[이미지가 설명하는 배경·감정]\n\n[이모지 소제목] 복선\n[사실 + 실제 장면]\n\n[네이버 구분선]\n\n[이모지 소제목] 전환과 결과\n[제목 약속 회수]\n\n[이모지 소제목] 브링이슈의 한 줄 해석\n[독자에게 남기는 판단]\n\n[댓글 질문 1개]\n\n출처: [채널 / 영상 / URL / 확인일]", LIGHT_GRAY)

    add_heading(doc, "8. 네이버 편집기 서식 매뉴얼", 1)
    add_heading(doc, "8.1 기본 서식", 2)
    add_table(doc, ["요소", "브링이슈 기준"], [
        ("정렬", "가운데 정렬 기본"),
        ("글꼴", "나눔고딕"),
        ("본문 크기", "네이버 편집기 19"),
        ("줄간격", "180%"),
        ("문단", "1~3문장, 의미 묶음 사이 빈 문단 1개"),
        ("소제목", "이모지 + 짧은 문장 + 연두/노랑 하이라이트"),
        ("인용구", "전환 또는 핵심 판단 1회"),
        ("밑줄", "가장 중요한 반전·결론 1~2문장만"),
        ("굵게", "핵심 명사·결과 3~5곳"),
        ("구분선", "정보 블록이 바뀌는 지점 1~3회"),
        ("이미지", "문서 폭에 맞춤, 장면 가까이에 분산 배치"),
        ("캡션", "크기 15, 회색, 가운데 정렬"),
        ("이모지", "글 전체 5~8개"),
        ("태그", "관련성 높은 6~10개"),
    ], [1875, 7485], font_size=9.5)
    add_heading(doc, "8.2 이모지·하이라이트 규칙", 2)
    add_bullet(doc, "🔥 결과·화제성 / 👀 관찰·복선 / 🎬 장면 전환 / 💬 대화 / ✅ 결론 / 📌 저장 정보처럼 역할을 정해 쓴다.", bullet_num)
    add_bullet(doc, "모든 문장에 이모지를 붙이지 않는다. 소제목과 감정 전환에만 사용한다.", bullet_num)
    add_bullet(doc, "하이라이트는 소제목 또는 한 문장 일부에만 사용한다. 문단 전체를 칠하지 않는다.", bullet_num)
    add_bullet(doc, "밑줄과 굵게와 하이라이트를 한 문장에 동시에 겹치지 않는다.", bullet_num)
    add_heading(doc, "8.3 모바일 리듬 검수", 2)
    add_checklist(doc, [
        "첫 화면에서 인물·사건·궁금증이 보인다.",
        "긴 문단이 4줄 이상 연속되지 않는다.",
        "사진과 캡션이 설명 문단 바로 가까이에 있다.",
        "구분선이 너무 자주 등장해 글이 조각나지 않는다.",
        "같은 문장 구조와 키워드가 반복되지 않는다.",
        "출처 블록이 본문보다 과도하게 커 보이지 않는다.",
    ])

    add_heading(doc, "9. 네이버 업로드·발행 SOP", 1)
    add_callout(doc, "중요", "원고·이미지 검수가 끝나도 바로 공개하지 않습니다. 전체 미리보기를 보여주고 사용자의 최종 발행 승인을 받은 뒤에만 공개합니다.", PALE_RED)
    steps = [
        "브링이슈 계정과 블로그 주소(blog.naver.com/bringissue)를 확인한다.",
        "새 글 쓰기를 열고 카테고리를 ‘오늘 뜬 유튜브’로 선택한다.",
        "제목을 입력하고 모바일 앞부분에서 인물·사건이 잘리는지 확인한다.",
        "본문 전체를 붙여 넣은 뒤 기본 정렬·글꼴·크기·줄간격을 적용한다.",
        "의미 묶음 사이에 네이버 편집기 빈 문단 1개를 만든다.",
        "소제목마다 이모지와 승인된 하이라이트 색을 적용한다.",
        "핵심 판단 1곳에 인용구 컴포넌트를 적용한다.",
        "반전·결론 1~2문장에만 밑줄을 적용한다.",
        "핵심 명사·결과 3~5곳만 굵게 처리한다.",
        "장면표 순서대로 깨끗한 16:9 캡처를 업로드한다.",
        "각 사진을 관련 문단 바로 아래에 배치하고 캡션을 입력한다.",
        "문장 전용 정보 카드 대신 승인된 제미나이 매거진 일러스트를 해설 문단에 배치한다.",
        "모든 제미나이 이미지 아래에 ‘AI로 제작한 이해용 이미지입니다.’ 캡션을 입력한다.",
        "정보 블록이 바뀌는 지점에 실제 구분선 컴포넌트를 넣는다.",
        "원본 영상 링크를 출처에 넣고 채널·영상 제목을 표시한다.",
        "태그 6~10개를 입력하고 관련 없는 인기 태그는 제거한다.",
        "주제는 ‘스타·연예인’ 등 실제 내용과 맞는 항목으로 설정한다.",
        "전체공개·검색허용·댓글·공감 설정을 확인한다.",
        "대표 이미지를 클라이맥스 장면으로 지정한다.",
        "모바일 미리보기에서 문단, 캡션, 이미지 비율, 링크를 확인한다.",
        "오탈자·중복문장·AI 티가 나는 상투 표현을 제거한다.",
        "개인정보 가림과 출처·저작권 기준을 다시 확인한다.",
        "원고 전체와 대표 이미지를 사용자에게 보여준다.",
        "사용자의 명시적인 최종 발행 승인을 받는다.",
        "승인 후 공개 발행한다.",
        "공개된 글을 다시 열어 실제 화면과 링크를 확인한다.",
        "발행 시각과 URL을 성과 원장에 기록한다.",
    ]
    decimal_num = decimal_nums[1]
    for step in steps:
        add_number(doc, step, decimal_num)

    doc.add_page_break()
    add_heading(doc, "10. 사실·권리·개인정보 검수", 1)
    add_heading(doc, "10.1 사실을 세 층으로 나눈다", 2)
    add_table(doc, ["구분", "표현 방식", "예시"], [
        ("영상에서 확인", "‘영상에서는 ~가 보였다’", "행동, 짧은 대화, 장소"),
        ("공식 출처 확인", "출처와 확인일을 기록", "제품명, 가격, 일정, 소속"),
        ("브링이슈 해석", "‘우리가 보기에는’, ‘이 장면의 포인트는’", "성격·선택·반응의 의미"),
    ], [1800, 3500, 4060], font_size=9.3)
    add_heading(doc, "10.2 영상 캡처와 인용 원칙", 2)
    add_bullet(doc, "원본을 대체하지 않는 범위에서 비평·설명에 필요한 장면만 사용한다.", bullet_num)
    add_bullet(doc, "대사 전문·연속 스크린샷·긴 클립으로 원본의 핵심 소비를 대신하지 않는다.", bullet_num)
    add_bullet(doc, "채널 워터마크를 제거하거나 다른 출처처럼 보이게 편집하지 않는다.", bullet_num)
    add_bullet(doc, "원본 영상 링크와 출처를 독자가 확인할 수 있게 제공한다.", bullet_num)
    add_bullet(doc, "권리 문제가 명확하지 않으면 해당 장면을 빼거나 자체 제작 설명 이미지로 대체한다.", bullet_num)
    add_bullet(doc, "제미나이 이미지는 실제 장면이나 보도사진처럼 오해되지 않도록 에디토리얼 일러스트로 제작하고 AI 이미지임을 표시한다.", bullet_num)
    add_heading(doc, "10.3 공개 전 강제 체크", 2)
    add_checklist(doc, [
        "영상 속 주장을 객관적 사실처럼 바꾸지 않았다.",
        "정책·가격·제품·일정 등 변동 정보는 발행 당일 공식 출처로 확인했다.",
        "얼굴·전화번호·계정·차량번호·위치·문서 식별정보를 점검했다.",
        "미성년자 또는 일반인의 사생활을 불필요하게 확대하지 않았다.",
        "제휴 링크 앞에 광고 고지가 있다.",
        "제목의 강한 표현을 본문 근거로 회수했다.",
    ])

    doc.add_page_break()
    add_heading(doc, "11. 수익화 연결 매뉴얼", 1)
    add_heading(doc, "11.1 수익화 판단 순서", 2)
    decimal_num = decimal_nums[2]
    for text in [
        "영상 안에서 실제 구매 문제나 제품 관심이 발생했는지 확인한다.",
        "공식 브랜드·제작진·착장 정보 등으로 제품을 검증한다.",
        "독자가 비교할 기준과 주의점을 먼저 정리한다.",
        "승인된 제휴 링크 또는 판매 경로가 있는지 확인한다.",
        "본편과 분리한 후속 글로 발행하고, 필요한 경우 본편에서 관련글로 연결한다.",
        "조회수와 구매 클릭·전환을 별도 지표로 기록한다.",
    ]:
        add_number(doc, text, decimal_num)
    add_heading(doc, "11.2 연결 가능한 후속 글", 2)
    add_bullet(doc, "영상 속 공식 확인 제품·착장·뷰티 아이템 정리", bullet_num)
    add_bullet(doc, "인물이 사용한 방법을 독자가 실행할 수 있게 바꾼 가이드", bullet_num)
    add_bullet(doc, "가격대·대체재·사용 상황을 비교한 구매 판단 글", bullet_num)
    add_bullet(doc, "영상 속 돈·계약·선택을 일반 독자의 의사결정으로 확장한 정보 글", bullet_num)
    add_callout(doc, "금지", "상품명이 확인되지 않았는데 ‘고윤정이 쓴 제품’, ‘OO가 선택한 상품’처럼 단정하지 않습니다. 유사 상품 추천은 인물과의 직접 연관을 제거하고 별도 비교 글로 작성합니다.", PALE_RED)

    doc.add_page_break()
    add_heading(doc, "12. 발행 후 성과 측정과 개선", 1)
    add_heading(doc, "12.1 측정 시점", 2)
    add_table(doc, ["시점", "확인 항목", "결정"], [
        ("1시간", "초기 조회·반응·표시 오류", "제목·대표 이미지의 즉시 문제만 확인"),
        ("6시간", "조회 증가 속도·유입 경로", "홈피드/검색/이웃 반응 가설 기록"),
        ("24시간", "조회·공감·댓글·체류 신호", "후속 글 여부 판단"),
        ("72시간", "유입 경로·검색어·대표 이미지 반응", "제목·도입 패턴 학습"),
        ("7일", "검색 유입·관련글 이동·독자 구성", "카테고리와 내부 연결 보완"),
        ("30일", "누적 조회·검색 지속성·구매 전환", "다음 달 포트폴리오 조정"),
    ], [1200, 4100, 4060], font_size=9.2)
    add_heading(doc, "12.2 성과 원장", 2)
    add_table(doc, ["글 URL", "발행일", "주제/각도", "1h", "24h", "72h", "7d", "30d", "주요 유입", "전환"], [
        ("", "", "", "", "", "", "", "", "", ""),
        ("", "", "", "", "", "", "", "", "", ""),
        ("", "", "", "", "", "", "", "", "", ""),
    ], [1100, 720, 1350, 600, 600, 600, 600, 650, 1700, 1440], font_size=7.8)
    add_heading(doc, "12.3 데이터 해석 원칙", 2)
    add_bullet(doc, "‘일반서비스’ 유입을 전부 홈피드라고 단정하지 않는다.", bullet_num)
    add_bullet(doc, "조회수와 구매·문의 전환은 따로 본다.", bullet_num)
    add_bullet(doc, "한 글의 성과로 규칙을 만들지 않고 최소 10개 글의 패턴을 비교한다.", bullet_num)
    add_bullet(doc, "제목, 대표 이미지, 도입, 주제, 발행 시각 중 무엇이 영향을 줬는지 가설을 하나씩 분리한다.", bullet_num)
    add_bullet(doc, "확인할 수 없는 값은 0으로 만들지 않고 NA로 기록한다.", bullet_num)

    add_heading(doc, "13. 고윤정 영상 적용 예시", 1)
    add_para(doc, "대상 영상: ‘고윤정을 좋아하세요?’ / 안녕하세요원이입니다잘부탁드립니다 / 약 19분. 아래는 템플릿 적용 예시이며, 변동 조회수는 발행 시점에 다시 확인합니다.")
    add_heading(doc, "13.1 기획 결론", 2)
    add_table(doc, ["항목", "적용 내용"], [
        ("본편 각도", "팬심이 실제 만남으로 이어진 과정"),
        ("핵심 질문", "좋아하는 마음은 어떤 행동을 만들었고, 그 행동은 어떻게 현실이 되었나?"),
        ("제목 약속", "학교 방문 → 예상 밖 만남 → 식사 제안 → 연락처 교환"),
        ("브링이슈 해설", "우연만이 아니라 행동의 누적이 결과를 만든 이야기"),
        ("후속 수익 글", "공식 확인 가능한 고윤정 뷰티·패션 정보만 별도 제작"),
        ("댓글 질문", "여러분이라면 좋아하는 사람을 만나기 위해 어디까지 해볼 수 있나요?"),
    ], [1875, 7485], font_size=9.4)
    add_heading(doc, "13.2 장면표", 2)
    add_table(doc, ["번호", "타임코드", "본문 역할", "장면"], [
        ("1", "00:51", "도입", "팬심의 출발"),
        ("2", "03:00", "인물 취향", "중저음 목소리 포인트"),
        ("3", "04:41", "행동", "서울여대 도착"),
        ("4", "06:20", "복선", "조형예술관 이동"),
        ("5", "12:24", "전환", "고윤정이 위층에 있다는 사실을 알게 됨"),
        ("6", "13:38", "대표", "첫 만남"),
        ("7", "14:28", "대화", "닮았다는 반응"),
        ("8", "15:51", "클라이맥스", "식사 제안"),
        ("9", "16:32", "결과", "연락처 교환"),
        ("10", "17:34", "여운", "만남 이후 감정 반응"),
    ], [700, 1150, 1700, 5810], font_size=9)
    add_heading(doc, "13.3 제목 후보", 2)
    decimal_num = decimal_nums[3]
    add_number(doc, "고윤정 팬이 서울여대까지 찾아간 이유, 결국 연락처까지 받았다", decimal_num)
    add_number(doc, "‘고윤정을 좋아하세요?’ 팬심 하나로 시작해 실제 만남까지 간 과정", decimal_num)
    add_number(doc, "학교에 갔을 뿐인데 고윤정이 위층에 있었다… 그 뒤 벌어진 일", decimal_num)
    add_callout(doc, "추천 방향", "1번처럼 인물·행동·결과가 보이는 사건형을 우선합니다. 다만 발행 직전 전체 영상과 실제 장면을 다시 확인해 표현을 최종 확정합니다.", PALE_YELLOW)

    add_heading(doc, "부록 A. 1편 제작용 빈 브리프", 1)
    for label, hint in [
        ("작업일 / 담당", ""), ("영상 URL", ""), ("채널 / 영상 제목", ""),
        ("확인 시점 지표", "조회수 / 좋아요 / 댓글 / 확인 시각"),
        ("한 줄 사건", ""), ("예상 독자", ""), ("독자의 한 질문", ""),
        ("이 글의 한 약속", ""), ("본편 각도", ""), ("브링이슈의 새 해석", ""),
        ("대표 이미지 타임코드", ""), ("댓글 질문", ""), ("후속 수익 글", "없음 / 검증 후 제작"),
        ("추천 제목", ""), ("대안 제목 1", ""), ("대안 제목 2", ""),
    ]:
        add_form_line(doc, label, hint)
    add_heading(doc, "부록 B. 최종 발행 승인표", 1)
    add_checklist(doc, [
        "영상 전체를 확인했고 대사·타임라인 원장이 있다.",
        "75점 이상이며 강제 제외 조건이 없다.",
        "제목의 약속을 도입과 본문에서 회수했다.",
        "브링이슈만의 해석이 단순 요약보다 분명하다.",
        "캡처가 모두 장면표에 있고 YouTube UX/UI가 제거됐다.",
        "문장 전용 정보 카드가 없고 제미나이 이미지는 배경·감정·맥락 설명에만 사용됐다.",
        "모든 AI 이미지에 ‘AI로 제작한 이해용 이미지입니다.’ 캡션이 있다.",
        "제미나이 이미지가 1200×675px 이상이며 얼굴 복제·깨진 글자·로고·워터마크가 없다.",
        "원본 워터마크를 지우지 않았고 개인정보를 가렸다.",
        "문단·이모지·하이라이트·밑줄·인용구·구분선 기준을 지켰다.",
        "공식 확인이 필요한 변동 정보를 발행 당일 확인했다.",
        "제휴 링크가 있으면 광고 고지를 넣었다.",
        "모바일 미리보기와 링크를 점검했다.",
        "사용자에게 최종 원고와 대표 이미지를 보여줬다.",
        "사용자가 공개 발행을 명시적으로 승인했다.",
    ])
    add_form_line(doc, "최종 승인자", "")
    add_form_line(doc, "승인 일시", "")
    add_form_line(doc, "발행 URL", "")

    add_heading(doc, "부록 C. 공식 참고 기준", 1)
    add_para(doc, "네이버가 공개하지 않은 글자 수·사진 수·발행 시각·태그 수·키워드 반복 횟수를 홈피드 공식 공식처럼 사용하지 않습니다. 아래 공개 안내와 실제 블로그 통계를 기준으로 판단합니다.")
    sources = [
        ("네이버 홈피드 소개", "https://help.naver.com/service/5630/contents/23396?lang=ko&osType=MOBILE"),
        ("네이버 콘텐츠 작성 권장 사항", "https://searchadvisor.naver.com/guide/content-basic"),
        ("네이버 블로그 검색 노출 순서", "https://help.naver.com/service/5593/contents/10582?lang=ko&osType=COMMONOS"),
        ("원본 영상 예시", "https://www.youtube.com/watch?v=6y_g78FQ3io"),
    ]
    for title, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(title + ": ")
        set_run_font(r1, 10.5, True, DARK_BLUE)
        r2 = p.add_run(url)
        set_run_font(r2, 9.5, False, MID_BLUE, underline=True)
    add_callout(doc, "매뉴얼 개정 규칙", "새로운 운영 규칙은 ‘느낌’이 아니라 최소 10개 글의 실제 데이터, 네이버 공식 안내, 권리·정책 변경 중 하나를 근거로 추가합니다. 개정 시 버전·날짜·변경 이유를 문서 첫머리에 남깁니다.", PALE_GREEN)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build()
